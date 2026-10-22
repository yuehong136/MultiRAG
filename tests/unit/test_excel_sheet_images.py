"""Real XLSX drawings must stay attached to their source worksheet."""

from io import BytesIO
from typing import Any

import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as WorksheetImage
from PIL import Image

from core.app import table
from core.utils.lazy_image import ensure_pil_image


@pytest.mark.parametrize("second_image", [False, True])
@pytest.mark.parametrize("description_list", [False, True])
def test_cell_images_do_not_leak_to_later_sheets(monkeypatch: pytest.MonkeyPatch, second_image: bool, description_list: bool) -> None:
    book = Workbook()
    first = book.active
    assert first is not None
    first.title = "First"
    second = book.create_sheet("Second")
    third = book.create_sheet("Third")
    for sheet in (first, second, third):
        sheet.append(["name", "picture"])
        sheet.append([sheet.title, None])

    for sheet, color, anchor in ((first, "red", "B2"), (first, "green", "D2"), (second, "blue", "B2")):
        if sheet is second and not second_image:
            continue
        png = BytesIO()
        Image.new("RGB", (32, 32), color=color).save(png, format="PNG")
        sheet.add_image(WorksheetImage(png), anchor)

    source = BytesIO()
    book.save(source)
    described: list[tuple[str, tuple[int, ...]]] = []

    def describe(images: list[dict[str, Any]], **kwargs: Any) -> list[Any]:
        results = []
        for image in images:
            pixels = ensure_pil_image(image["image"])
            described.append((image["sheet"], pixels.getpixel((0, 0))))
            description = f"{image['sheet']} column {image['col_from']}"
            results.append(((image["image"], [description] if description_list else description), [(0, 0, 0, 0, 0)]))
        return results

    monkeypatch.setattr(table, "vision_figure_parser_figure_xlsx_wrapper", describe)
    frames, figures = table.Excel()("source.xlsx", binary=source.getvalue(), callback=lambda *args: None)

    assert [frame["name"].tolist() for frame in frames] == [["First"], ["Second"], ["Third"]]
    assert frames[0].iat[0, 1] == "First column 2"
    assert frames[1].iat[0, 1] == ("Second column 2" if second_image else None)
    assert frames[2].iat[0, 1] is None
    assert [figure[0][1] for figure in figures] == [["First column 4"]]
    assert described == [("First", (255, 0, 0)), ("First", (0, 128, 0))] + ([("Second", (0, 0, 255))] if second_image else [])
