"""Real XLSX boundaries consumed by the HTML ingestion paths."""

from io import BytesIO

import pytest
from bs4 import BeautifulSoup
from openpyxl import Workbook

from core.app import naive, one
from deepdoc.parser.excel_parser import RAGFlowExcelParser


def workbook_bytes(count: int, *, extra_sheets: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Rows"
    sheet.append(["name", "number", "flag"])
    for number in range(count):
        sheet.append([f"row <{number}>", number, False])
    # Formatted blank rows must not create a trailing chunk.
    sheet.cell(row=count + 20, column=1).number_format = "0.00"
    if extra_sheets:
        workbook.create_sheet("Empty")
        workbook.create_sheet("Header").append(["only header"])
        second = workbook.create_sheet("Second")
        second.append(["value"])
        second.append(["last sheet"])
    binary = BytesIO()
    workbook.save(binary)
    return binary.getvalue()


@pytest.mark.parametrize("count, sizes", [(0, []), (1, [1]), (12, [12]), (13, [12, 1]), (24, [12, 12]), (25, [12, 12, 1])])
def test_html_chunks_keep_all_data_and_emit_no_header_only_tail(count: int, sizes: list[int]) -> None:
    chunks = RAGFlowExcelParser().html(workbook_bytes(count), chunk_rows=12)
    tables = [BeautifulSoup(chunk, "html.parser") for chunk in chunks]
    assert [len(table.select("tr")) - 1 for table in tables] == sizes
    assert all([cell.text for cell in table.select("th")] == ["name", "number", "flag"] for table in tables)
    cells = [row.select("td") for table in tables for row in table.select("tr")[1:]]
    assert [[cell.text for cell in row] for row in cells] == [[f"row <{number}>", str(number), "False"] for number in range(count)]
    assert all("&lt;" in chunk for chunk in chunks)


def test_html_sheet_boundaries_do_not_emit_empty_or_header_only_sheets() -> None:
    chunks = RAGFlowExcelParser().html(workbook_bytes(12, extra_sheets=True), chunk_rows=12)
    assert [BeautifulSoup(chunk, "html.parser").caption.text for chunk in chunks] == ["Rows", "Second"]
    assert "last sheet" in chunks[-1]


@pytest.mark.parametrize("chunk_rows", [0, -1])
def test_html_rejects_invalid_chunk_size(chunk_rows: int) -> None:
    with pytest.raises(ValueError, match="positive"):
        RAGFlowExcelParser().html(workbook_bytes(1), chunk_rows=chunk_rows)


@pytest.mark.parametrize("consumer", ["naive", "one"])
def test_real_html_ingestion_consumers_keep_exact_multiple_chunks(monkeypatch: pytest.MonkeyPatch, consumer: str) -> None:
    binary = workbook_bytes(24)
    if consumer == "naive":
        chunks = naive.chunk("rows.xlsx", binary=binary, callback=lambda *a: None, parser_config={"html4excel": True, "chunk_token_num": 12})
        assert len(chunks) == 2
    else:
        monkeypatch.setattr(one, "MAXIMUM_TASK_PAGE_NUMBER", 12)
        chunks = one.chunk("rows.xlsx", binary=binary, callback=lambda *a: None)
        assert len(chunks) == 1
    content = "\n".join(chunk["content_with_weight"] for chunk in chunks)
    tables = BeautifulSoup(content, "html.parser").select("table")
    assert [len(table.select("tr")) - 1 for table in tables] == [12, 12]
