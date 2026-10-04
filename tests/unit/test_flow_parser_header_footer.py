import json
from io import BytesIO
from types import SimpleNamespace
from typing import Any

import pytest
from docx import Document

from core.flow.parser.header_footer import remove_header_footer_docx_blob
from core.flow.parser.parser import Parser, ParserParam
from core.flow.parser.pdf_chunk_metadata import PDF_POSITIONS_KEY


def make_parser(file_type: str, enabled: bool | str, output: str = "json") -> Parser:
    param = ParserParam()
    param.setups[file_type].update(remove_header_footer=enabled, output_format=output)
    parser = Parser.__new__(Parser)
    parser._param = param
    parser._canvas = SimpleNamespace(_tenant_id="tenant-1")
    parser.callback = lambda *args, **kwargs: None
    parser._id = "parser-1"
    return parser


@pytest.mark.parametrize("enabled", [True, False, "true", "false"])
def test_flags_save_reload_as_booleans(enabled: bool | str) -> None:
    param = ParserParam()
    for kind in ("pdf", "doc", "docx", "html"):
        assert param.setups[kind]["remove_header_footer"] is False
        param.setups[kind].update(remove_header_footer=enabled, remove_toc=enabled)
    param.setups = {kind: param.setups[kind] for kind in ("pdf", "doc", "docx", "html")}
    param.check()
    reloaded = ParserParam()
    reloaded.update(json.loads(json.dumps(param.as_dict())))
    reloaded.check()
    for kind in ("pdf", "doc", "docx", "html"):
        assert reloaded.setups[kind]["remove_header_footer"] is (enabled in (True, "true"))
        assert reloaded.setups[kind]["remove_toc"] is (enabled in (True, "true"))


@pytest.mark.parametrize("enabled", [True, False, "false"])
@pytest.mark.parametrize("output", ["json", "text"])
def test_html_filters_structure_preserves_body_and_fragment(enabled: bool | str, output: str) -> None:
    parser = make_parser("html", enabled, output)
    parser._html(
        "fragment.html",
        b'<header>masthead</header><div role="contentinfo"><footer>footer-only</footer></div><main><p>masthead</p><p>Body &amp; detail</p><table><tr><th>Column</th></tr><tr><td>Value</td></tr></table></main>',
    )
    result = str(parser.output(output))
    assert "Body &amp; detail" in result or "Body & detail" in result
    assert "masthead" in result
    assert "Column" in result and "Value" in result
    assert ("footer-only" not in result) is (enabled is True)


@pytest.mark.parametrize("remove_toc", [True, False, "true", "false"])
def test_html_toc_boolean_contract(monkeypatch: pytest.MonkeyPatch, remove_toc: bool | str) -> None:
    import core.flow.parser.parser as module

    calls: list[Any] = []

    def filter_toc(items: list[Any]) -> tuple[list[Any], list[int]]:
        calls.append(items)
        return items, list(range(len(items)))

    monkeypatch.setattr(module, "remove_toc", filter_toc)
    parser = make_parser("html", False)
    parser._param.setups["html"]["remove_toc"] = remove_toc
    parser._html("body.html", b"<p>Body</p>")
    assert bool(calls) is (remove_toc in (True, "true"))


def docx_blob() -> bytes:
    document = Document()
    for container in (document.sections[0].header, document.sections[0].footer, document.sections[0].first_page_header, document.sections[0].even_page_footer):
        container.paragraphs[0].text = "repeated body text"
        container.add_paragraph("header-footer-only")
    document.add_paragraph("repeated body text")
    document.add_paragraph("Body paragraph")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Body table"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_docx_structure_removal_keeps_identical_body_text_and_tables() -> None:
    original = docx_blob()
    filtered = Document(BytesIO(remove_header_footer_docx_blob(original)))
    assert [p.text for p in filtered.paragraphs] == ["repeated body text", "Body paragraph"]
    assert filtered.tables[0].cell(0, 0).text == "Body table"
    assert all(not p.text for p in filtered.sections[0].header.paragraphs)
    assert all(not p.text for p in filtered.sections[0].even_page_footer.paragraphs)


@pytest.mark.parametrize("enabled", [True, False])
@pytest.mark.parametrize("output", ["json", "markdown"])
def test_docx_real_parser_preserves_body(enabled: bool, output: str) -> None:
    parser = make_parser("docx", enabled, output)
    parser._docx("body.docx", docx_blob())
    result = str(parser.output(output))
    assert "repeated body text" in result and "Body paragraph" in result and "Body table" in result
    # Both existing DOCX engines already traverse only the document body.
    assert "header-footer-only" not in result


@pytest.mark.parametrize("method", ["_doc", "_docx"])
@pytest.mark.parametrize("enabled", [True, False])
@pytest.mark.parametrize("output", ["json", "markdown"])
def test_doc_tika_structural_filter(monkeypatch: pytest.MonkeyPatch, method: str, enabled: bool, output: str) -> None:
    from tika import parser as tika_parser

    def parse(_stream: BytesIO, **kwargs: Any) -> dict[str, str]:
        assert kwargs == ({"xmlContent": True} if enabled else {})
        return {
            "content": '<html><body><div class="header"><p>repeated</p></div><p>repeated</p><p>Body</p><div class="footer"><p>footer-only</p></div></body></html>'
            if enabled
            else "repeated\nrepeated\nBody\nfooter-only"
        }

    monkeypatch.setattr(tika_parser, "from_buffer", parse)
    parser = make_parser(method[1:], enabled, output)
    getattr(parser, method)("body.doc", b"doc")
    result = str(parser.output(output))
    assert "repeated" in result and "Body" in result
    assert ("footer-only" not in result) is enabled


@pytest.mark.parametrize("enabled", [True, False, "false"])
@pytest.mark.parametrize("output", ["json", "markdown"])
def test_pdf_preserves_metadata_author_abstract_and_table_headers(monkeypatch: pytest.MonkeyPatch, enabled: bool | str, output: str) -> None:
    import core.flow.parser.parser as module

    class FakePdfParser:
        outlines: list[Any] = []

        def parse_into_bboxes(self, _blob: bytes, callback: Any) -> list[dict[str, Any]]:
            rows = [
                ("header-only", "header"),
                ("Research title", "title"),
                ("Alice", "text"),
                ("Abstract", "title"),
                ("body " * 40, "text"),
                ("Column", "table header"),
                ("footer-only", "page_footer"),
                ("number-only", "page number"),
            ]
            return [{"text": text, "layout_type": layout, "page_number": 3, "position_tag": "@@3\t10\t30\t20\t40##"} for text, layout in rows]

    monkeypatch.setattr(module, "RAGFlowPdfParser", FakePdfParser)
    parser = make_parser("pdf", enabled, output)
    parser._param.setups["pdf"]["preprocess"] = ["author", "abstract"]
    parser._pdf("body.pdf", b"pdf")
    result = str(parser.output(output))
    assert "Column" in result and "Alice" in result and "body " in result
    for removed in ("header-only", "footer-only", "number-only"):
        assert (removed not in result) is (enabled is True)
    if output == "json":
        rows = parser.output(output)
        assert all(row["layoutno"] for row in rows)
        assert all(row[PDF_POSITIONS_KEY] == [[3, 10.0, 30.0, 20.0, 40.0]] for row in rows)
        assert next(row for row in rows if row["text"] == "Alice")["author"] is True
        assert next(row for row in rows if row["text"].startswith("body "))["abstract"] is True
